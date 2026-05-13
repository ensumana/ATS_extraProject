"""
ATS Backend - No Job Description + Duplicate Resume Detection
--------------------------------------------------------------
Ready-to-replace backend for the Streamlit ATS app.

What changed from the earlier JD-based version:
1. Job description is no longer required.
2. Candidate scoring is based on skill match + experience.
3. Duplicate/repeated resumes are detected within the current uploaded batch.
4. If a duplicate is found, the row is still displayed with Duplicate Status = "Candidate Exists".
5. Backward-compatible function names are preserved where possible.

Required packages:
    pip install pdfplumber python-docx pandas openpyxl

Optional:
    sentence-transformers is NOT required for this version.
"""

from __future__ import annotations

import hashlib
import logging
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import docx
import pandas as pd
import pdfplumber

# -----------------------------------------------------------------------------
# Logging
# -----------------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)
logger = logging.getLogger("ats_backend")

# -----------------------------------------------------------------------------
# Constants
# -----------------------------------------------------------------------------
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_SIZE_MB = 15

EDUCATION_PATTERNS = [
    ("Ph.D", [r"\bph\.?d\.?\b", r"doctor of philosophy"]),
    ("M.Tech", [r"\bm\.?tech\b", r"master of technology"]),
    ("M.E", [r"\bm\.?e\.?\b", r"master of engineering"]),
    ("M.Sc", [r"\bm\.?sc\b", r"master of science"]),
    ("MBA", [r"\bm\.?b\.?a\.?\b", r"master of business administration"]),
    ("B.Tech", [r"\bb\.?tech\b", r"bachelor of technology"]),
    ("B.E", [r"\bb\.?e\.?\b", r"bachelor of engineering"]),
    ("B.Sc", [r"\bb\.?sc\b", r"bachelor of science"]),
    ("Diploma", [r"\bdiploma\b"]),
]

COMMON_HEADER_WORDS = {
    "resume", "curriculum vitae", "cv", "profile", "career objective", "summary",
    "professional summary", "contact", "personal details", "education", "skills"
}

# -----------------------------------------------------------------------------
# Result container
# -----------------------------------------------------------------------------
@dataclass
class ResumeResult:
    candidate_name: str
    file_name: str
    email: str
    phone: str
    education: str
    experience_years: float
    skills: str
    matched_skill_count: int
    total_required_skills: int
    skill_score: float
    experience_score: float
    final_score: float
    category: str
    duplicate_status: str
    duplicate_reason: str
    existing_candidate: str
    why_selected: str
    error: str = ""

    def to_display_dict(self) -> Dict[str, object]:
        return {
            "Candidate Name": self.candidate_name,
            "File Name": self.file_name,
            "Email": self.email,
            "Phone": self.phone,
            "Education": self.education,
            "Experience": self.experience_years,
            "Skills": self.skills,
            "Matched Skills": self.matched_skill_count,
            "Required Skills": self.total_required_skills,
            "Skill Score": self.skill_score,
            "Experience Score": self.experience_score,
            "Score": self.final_score,
            "Category": self.category,
            "Duplicate Status": self.duplicate_status,
            "Duplicate Reason": self.duplicate_reason,
            "Existing Candidate": self.existing_candidate,
            "Why Selected": self.why_selected,
            "Error": self.error,
        }

# -----------------------------------------------------------------------------
# Backend engine
# -----------------------------------------------------------------------------
class ATSBackend:
    def __init__(self):
        self.skills: List[str] = []
        self.min_required_experience: float = 0.0
        self.results: List[Dict[str, object]] = []
        self._seen_candidates: Dict[str, Dict[str, str]] = {}

    # -------------------------------------------------------------------------
    # Backward-compatible no-op. Kept so older app imports do not fail.
    # -------------------------------------------------------------------------
    def load_jd(self, jd_text: str) -> None:
        logger.info("JD input ignored. This backend version does not use job descriptions.")

    def set_skills(self, skills_str: str | Iterable[str]) -> None:
        if isinstance(skills_str, str):
            raw_skills = skills_str.split(",")
        else:
            raw_skills = list(skills_str)

        self.skills = sorted(
            {skill.strip().lower() for skill in raw_skills if skill and skill.strip()},
            key=len,
            reverse=True,
        )
        logger.info("Loaded %d required skills", len(self.skills))

    def set_min_required_experience(self, years: float) -> None:
        try:
            self.min_required_experience = max(float(years), 0.0)
        except (TypeError, ValueError):
            self.min_required_experience = 0.0

    def clear_results(self) -> None:
        self.results.clear()
        self._seen_candidates.clear()

    # -------------------------------------------------------------------------
    # File validation and extraction
    # -------------------------------------------------------------------------
    def validate_file(self, file_path: str | Path) -> Tuple[bool, str]:
        path = Path(file_path)
        if not path.exists():
            return False, "File does not exist"
        if not path.is_file():
            return False, "Path is not a file"
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            return False, f"Unsupported file type: {path.suffix}"
        size_mb = path.stat().st_size / (1024 * 1024)
        if size_mb > MAX_FILE_SIZE_MB:
            return False, f"File too large: {size_mb:.2f} MB"
        return True, ""

    def extract_text_original_case(self, file_path: str | Path) -> str:
        path = Path(file_path)
        valid, error = self.validate_file(path)
        if not valid:
            raise ValueError(error)

        text_parts: List[str] = []
        try:
            if path.suffix.lower() == ".pdf":
                with pdfplumber.open(path) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text() or ""
                        if page_text:
                            text_parts.append(page_text)
            elif path.suffix.lower() == ".docx":
                document = docx.Document(path)
                for para in document.paragraphs:
                    if para.text:
                        text_parts.append(para.text)
                for table in document.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text:
                                text_parts.append(cell.text)
        except Exception as exc:
            logger.exception("Failed to extract text from %s", path)
            raise RuntimeError(f"Failed to extract text: {exc}") from exc

        return "\n".join(text_parts)

    def extract_text(self, file_path: str | Path) -> str:
        return self.extract_text_original_case(file_path).lower()

    # -------------------------------------------------------------------------
    # Information extraction
    # -------------------------------------------------------------------------
    @staticmethod
    def extract_email(text: str) -> str:
        match = re.search(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}", text)
        return match.group(0).lower() if match else ""

    @staticmethod
    def extract_phone(text: str) -> str:
        phone_patterns = [
            r"(?:\+91[\s\-]?)?[6-9]\d{4}[\s\-]?\d{5}",
            r"(?:\+\d{1,3}[\s\-]?)?(?:\(?\d{2,4}\)?[\s\-]?)?\d{3,5}[\s\-]?\d{4,5}",
        ]
        for pattern in phone_patterns:
            match = re.search(pattern, text)
            if match:
                phone = match.group(0).strip()
                phone = re.sub(r"\s+", " ", phone)
                return phone
        return ""

    @staticmethod
    def normalize_phone(phone: str) -> str:
        digits = re.sub(r"\D", "", phone or "")
        if len(digits) > 10 and digits.endswith(digits[-10:]):
            # For Indian numbers, compare last 10 digits.
            return digits[-10:]
        return digits

    @staticmethod
    def extract_experience(text: str) -> float:
        candidates: List[float] = []
        patterns = [
            r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years|year|yrs|yr)\b",
            r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?)",
        ]
        for pattern in patterns:
            for match in re.findall(pattern, text, flags=re.IGNORECASE):
                try:
                    candidates.append(float(match))
                except ValueError:
                    pass

        detailed = re.findall(
            r"(\d+)\s*(?:years|year|yrs|yr)\s*(?:and)?\s*(\d+)\s*(?:months|month|mos|mo)",
            text,
            flags=re.IGNORECASE,
        )
        for years, months in detailed:
            candidates.append(float(years) + float(months) / 12.0)

        return round(max(candidates), 2) if candidates else 0.0

    @staticmethod
    def extract_education(text: str) -> str:
        found = []
        for label, patterns in EDUCATION_PATTERNS:
            if any(re.search(pattern, text, flags=re.IGNORECASE) for pattern in patterns):
                found.append(label)
        return ", ".join(found)

    @staticmethod
    def extract_candidate_name(original_text: str, file_path: str | Path) -> str:
        lines = [line.strip() for line in original_text.splitlines() if line.strip()]

        for line in lines[:12]:
            clean = re.sub(r"[^A-Za-z .'-]", "", line).strip()
            lower = clean.lower()
            words = clean.split()

            if not clean:
                continue
            if lower in COMMON_HEADER_WORDS:
                continue
            if any(word in lower for word in ["email", "phone", "mobile", "linkedin", "github"]):
                continue
            if 2 <= len(words) <= 4 and all(w[:1].isupper() for w in words if w):
                return clean

        return Path(file_path).stem.replace("_", " ").replace("-", " ").title()

    # -------------------------------------------------------------------------
    # Duplicate logic
    # -------------------------------------------------------------------------
    @staticmethod
    def content_hash(text: str) -> str:
        normalized = re.sub(r"\s+", " ", text.lower()).strip()
        return hashlib.sha256(normalized.encode("utf-8", errors="ignore")).hexdigest()

    def detect_duplicate(self, *, email: str, phone: str, candidate_name: str, text: str) -> Tuple[str, str, str]:
        """
        Return: duplicate_status, duplicate_reason, existing_candidate
        Duplicate priority:
        1. Same email
        2. Same phone
        3. Same resume content hash
        4. Same normalized candidate name, if no email/phone available
        """
        keys = []

        if email:
            keys.append((f"email:{email.lower()}", "Same email"))

        normalized_phone = self.normalize_phone(phone)
        if normalized_phone and len(normalized_phone) >= 10:
            keys.append((f"phone:{normalized_phone}", "Same phone number"))

        resume_hash = self.content_hash(text)
        if resume_hash:
            keys.append((f"hash:{resume_hash}", "Same resume content"))

        normalized_name = re.sub(r"\s+", " ", candidate_name.lower()).strip()
        if normalized_name and not email and not normalized_phone:
            keys.append((f"name:{normalized_name}", "Same candidate name"))

        for key, reason in keys:
            if key in self._seen_candidates:
                existing = self._seen_candidates[key]
                return "Candidate Exists", reason, existing.get("candidate_name", "")

        for key, _reason in keys:
            self._seen_candidates[key] = {
                "candidate_name": candidate_name,
                "email": email,
                "phone": phone,
            }

        return "New Candidate", "", ""

    # -------------------------------------------------------------------------
    # Scoring
    # -------------------------------------------------------------------------
    def skill_match(self, text: str) -> Tuple[int, str, float]:
        if not self.skills:
            return 0, "", 0.0

        found = []
        for skill in self.skills:
            pattern = rf"(?<!\w){re.escape(skill)}(?!\w)"
            if re.search(pattern, text, flags=re.IGNORECASE):
                found.append(skill)

        count = len(found)
        score = (count / len(self.skills)) * 100 if self.skills else 0.0
        return count, ", ".join(found), round(score, 2)

    def experience_score(self, experience_years: float) -> float:
        if self.min_required_experience <= 0:
            if experience_years <= 0:
                return 0.0
            if experience_years >= 5:
                return 100.0
            return round((experience_years / 5.0) * 100.0, 2)
        return round(min((experience_years / self.min_required_experience) * 100.0, 100.0), 2)

    @staticmethod
    def category_from_score(score: float, duplicate_status: str) -> str:
        if duplicate_status == "Candidate Exists":
            return "Duplicate"
        if score >= 75:
            return "Strong Match"
        if score >= 50:
            return "Moderate Match"
        return "Low Match"

    # -------------------------------------------------------------------------
    # Main processing
    # -------------------------------------------------------------------------
    def process_resume(self, file_path: str | Path) -> Dict[str, object]:
        path = Path(file_path)
        try:
            original_text = self.extract_text_original_case(path)
            text = original_text.lower()

            email = self.extract_email(original_text)
            phone = self.extract_phone(original_text)
            candidate_name = self.extract_candidate_name(original_text, path)
            experience = self.extract_experience(original_text)
            education = self.extract_education(original_text)
            skill_count, skills_found, skill_score = self.skill_match(text)
            exp_score = self.experience_score(experience)

            # No JD/semantic score in this version.
            # Weighted score: 80% skill fit + 20% experience fit.
            final_score = round((0.80 * skill_score) + (0.20 * exp_score), 2)

            duplicate_status, duplicate_reason, existing_candidate = self.detect_duplicate(
                email=email,
                phone=phone,
                candidate_name=candidate_name,
                text=text,
            )
            category = self.category_from_score(final_score, duplicate_status)

            if duplicate_status == "Candidate Exists":
                why_selected = (
                    f"Duplicate found: {duplicate_reason}. Existing candidate: {existing_candidate}. "
                    f"Matched skills: {skills_found or 'None'}, Experience: {experience} years."
                )
            else:
                why_selected = (
                    f"Matched {skill_count}/{len(self.skills)} required skills. "
                    f"Experience: {experience} years. Skill Score: {skill_score}, Experience Score: {exp_score}."
                )

            result = ResumeResult(
                candidate_name=candidate_name,
                file_name=path.name,
                email=email,
                phone=phone,
                education=education,
                experience_years=experience,
                skills=skills_found,
                matched_skill_count=skill_count,
                total_required_skills=len(self.skills),
                skill_score=skill_score,
                experience_score=exp_score,
                final_score=final_score,
                category=category,
                duplicate_status=duplicate_status,
                duplicate_reason=duplicate_reason,
                existing_candidate=existing_candidate,
                why_selected=why_selected,
            ).to_display_dict()

        except Exception as exc:
            logger.exception("Failed to process resume: %s", path)
            result = ResumeResult(
                candidate_name=path.stem,
                file_name=path.name,
                email="",
                phone="",
                education="",
                experience_years=0.0,
                skills="",
                matched_skill_count=0,
                total_required_skills=len(self.skills),
                skill_score=0.0,
                experience_score=0.0,
                final_score=0.0,
                category="Error",
                duplicate_status="Not Checked",
                duplicate_reason="",
                existing_candidate="",
                why_selected="Processing failed",
                error=str(exc),
            ).to_display_dict()

        self.results.append(result)
        return result

    def get_results_dataframe(self) -> pd.DataFrame:
        return pd.DataFrame(self.results)

    def process_folder(self, folder_path: str | Path) -> pd.DataFrame:
        folder = Path(folder_path)
        for file_path in folder.iterdir():
            if file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
                self.process_resume(file_path)
        return self.get_results_dataframe()

    def export_results(self, output_path: str | Path) -> str:
        df = self.get_results_dataframe()
        output = Path(output_path)
        if output.suffix.lower() == ".xlsx":
            df.to_excel(output, index=False)
        elif output.suffix.lower() == ".csv":
            df.to_csv(output, index=False)
        else:
            raise ValueError("Output path must end with .csv or .xlsx")
        return str(output)

# -----------------------------------------------------------------------------
# Backward-compatible module-level functions
# -----------------------------------------------------------------------------
_backend = ATSBackend()
results = _backend.results


def load_jd(jd_text: str) -> None:
    _backend.load_jd(jd_text)


def set_skills(skills_str: str | Iterable[str]) -> None:
    _backend.set_skills(skills_str)


def set_min_required_experience(years: float) -> None:
    _backend.set_min_required_experience(years)


def clear_results() -> None:
    _backend.clear_results()


def process_resume(file_path: str | Path) -> Dict[str, object]:
    return _backend.process_resume(file_path)


def process_folder(folder_path: str | Path) -> pd.DataFrame:
    return _backend.process_folder(folder_path)


def get_results_dataframe() -> pd.DataFrame:
    return _backend.get_results_dataframe()


def export_results(output_path: str | Path) -> str:
    return _backend.export_results(output_path)


def extract_text(file_path: str | Path) -> str:
    return _backend.extract_text(file_path)


def extract_email(text: str) -> str:
    return _backend.extract_email(text)


def extract_phone(text: str) -> str:
    return _backend.extract_phone(text)


def extract_experience(text: str) -> float:
    return _backend.extract_experience(text)


def extract_education(text: str) -> str:
    return _backend.extract_education(text)


def skill_score(text: str):
    count, skills_found, _score = _backend.skill_match(text)
    return count, skills_found
